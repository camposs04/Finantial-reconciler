# app/core/exporter.py
"""
Geração do arquivo Excel de saída com formatação profissional.

Usa openpyxl via pandas ExcelWriter para ter controle total sobre:
- Múltiplas abas
- Formatação de células (cores, fontes, bordas)
- Largura de colunas automática
- Aba de resumo executivo
"""

from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Optional

import pandas as pd
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter

from config.settings import (
    SHEET_CONCILIATED, SHEET_DIVERGENT,
    SHEET_MISSING_CLIENT, SHEET_MISSING_COMPANY,
    SHEET_DUPLICATES, SHEET_SUMMARY,
    COLOR_HEADER_BG, COLOR_HEADER_FG,
    COLOR_CONCILIATED, COLOR_DIVERGENT,
    COLOR_MISSING, COLOR_DUPLICATE,
)
from app.core.models import ReconciliationResult

logger = logging.getLogger(__name__)


def export_to_excel(result: ReconciliationResult) -> bytes:
    """
    Gera o arquivo Excel completo com todas as abas de resultado.
    
    Returns:
        bytes do arquivo Excel pronto para download
    """
    buffer = io.BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Ordem das abas: do mais importante ao mais específico
        _write_summary(writer, result)
        _write_sheet(
            writer, result.conciliated, SHEET_CONCILIATED,
            row_color=COLOR_CONCILIATED
        )
        _write_sheet(
            writer, result.divergent, SHEET_DIVERGENT,
            row_color=COLOR_DIVERGENT
        )
        _write_sheet(
            writer, result.missing_in_client, SHEET_MISSING_CLIENT,
            row_color=COLOR_MISSING
        )
        _write_sheet(
            writer, result.missing_in_company, SHEET_MISSING_COMPANY,
            row_color=COLOR_MISSING
        )
        _write_sheet(
            writer, result.duplicates, SHEET_DUPLICATES,
            row_color=COLOR_DUPLICATE
        )

    buffer.seek(0)
    return buffer.read()


# ── Funções privadas ────────────────────────────────────────────────────────────

def _write_sheet(
    writer: pd.ExcelWriter,
    df: pd.DataFrame,
    sheet_name: str,
    row_color: str = "FFFFFF",
) -> None:
    """
    Escreve um DataFrame em uma aba do Excel com formatação.
    Abas vazias recebem apenas uma mensagem informativa.
    """
    if df is None or df.empty:
        # Cria aba vazia com mensagem
        empty_df = pd.DataFrame({"Informação": ["Nenhum registro encontrado nesta categoria."]})
        empty_df.to_excel(writer, sheet_name=sheet_name, index=False)
        _format_sheet(writer.sheets[sheet_name], empty_df, row_color="FFFFFF")
        return

    # Remove colunas internas (prefixo _) do relatório
    display_df = df[[c for c in df.columns if not c.startswith("_")]].copy()
    display_df.to_excel(writer, sheet_name=sheet_name, index=False)
    _format_sheet(writer.sheets[sheet_name], display_df, row_color)


def _format_sheet(sheet, df: pd.DataFrame, row_color: str) -> None:
    """
    Aplica formatação profissional a uma aba do Excel:
    - Cabeçalho azul escuro com texto branco e negrito
    - Linhas de dados com cor de fundo conforme a categoria
    - Auto-ajuste da largura das colunas
    - Bordas finas em todas as células
    """
    header_fill = PatternFill("solid", start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG)
    header_font = Font(bold=True, color=COLOR_HEADER_FG, name="Arial", size=10)
    data_fill = PatternFill("solid", start_color=row_color, end_color=row_color)
    data_font = Font(name="Arial", size=9)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Formata cabeçalho (linha 1)
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = center

    # Formata linhas de dados
    for row_idx in range(2, sheet.max_row + 1):
        for col_idx in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.fill = data_fill
            cell.font = data_font
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    # Auto-ajuste de largura (limitado para não ficar largo demais)
    for col_idx, col_name in enumerate(df.columns, start=1):
        col_letter = get_column_letter(col_idx)
        # Calcula a largura baseada no conteúdo
        max_len = max(
            len(str(col_name)),
            *[len(str(sheet.cell(row=r, column=col_idx).value or ""))
              for r in range(2, min(sheet.max_row + 1, 52))],  # amostra primeiras 50 linhas
        )
        sheet.column_dimensions[col_letter].width = min(max_len + 2, 45)

    # Congela o cabeçalho (facilita navegação em tabelas grandes)
    sheet.freeze_panes = "A2"


def _write_summary(writer: pd.ExcelWriter, result: ReconciliationResult) -> None:
    """
    Gera a aba de resumo executivo — visão rápida dos resultados.
    """
    now = datetime.now().strftime("%d/%m/%Y %H:%M")

    summary_data = {
        "Métrica": [
            "📅 Data/Hora do Processamento",
            "── ENTRADAS ──",
            "Total de registros (Cliente)",
            "Total de registros (Empresa)",
            "── RESULTADOS ──",
            "✅ Conciliados",
            "⚠️  Divergentes (qtd. diferente)",
            "❌ Ausentes no Cliente",
            "❌ Ausentes na Empresa",
            "🔄 Duplicidades detectadas",
            "── INDICADOR ──",
            "Taxa de Conciliação",
        ],
        "Valor": [
            now,
            "",
            result.total_client_records,
            result.total_company_records,
            "",
            result.total_conciliated,
            result.total_divergent,
            result.total_missing_client,
            result.total_missing_company,
            result.total_duplicates,
            "",
            f"{result.reconciliation_rate}%",
        ],
    }

    summary_df = pd.DataFrame(summary_data)
    summary_df.to_excel(writer, sheet_name=SHEET_SUMMARY, index=False)

    sheet = writer.sheets[SHEET_SUMMARY]

    # Formatação especial para o resumo
    header_fill = PatternFill("solid", start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG)
    header_font = Font(bold=True, color=COLOR_HEADER_FG, name="Arial", size=10)

    for col_idx in range(1, 3):
        cell = sheet.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    # Destaca taxa de conciliação
    rate = result.reconciliation_rate
    rate_color = "C6EFCE" if rate >= 90 else ("FFEB9C" if rate >= 70 else "FFCCCC")

    for row_idx in range(2, sheet.max_row + 1):
        metric_cell = sheet.cell(row=row_idx, column=1)
        value_cell = sheet.cell(row=row_idx, column=2)
        
        is_rate_row = "Taxa" in str(metric_cell.value or "")
        fill_color = rate_color if is_rate_row else "F5F5F5"
        
        for cell in [metric_cell, value_cell]:
            cell.fill = PatternFill("solid", start_color=fill_color, end_color=fill_color)
            cell.font = Font(name="Arial", size=9, bold=is_rate_row)

    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 20
    sheet.freeze_panes = "A2"